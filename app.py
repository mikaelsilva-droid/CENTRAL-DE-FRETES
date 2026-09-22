import streamlit as st
import pandas as pd
import xml.etree.ElementTree as ET
import sqlite3
from datetime import datetime
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import matplotlib.pyplot as plt

# ReportLab para geração de PDFs
try:
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

st.set_page_config(page_title="Central de Fretes & Faturas", layout="wide")

# ==============================================================================
# ESTILIZAÇÃO CSS GLOBAL (TEMA AZUL E LINHAS INTERCALADAS DE ALTA VISIBILIDADE)
# ==============================================================================
st.markdown("""
<style>
    /* PALETA PRINCIPAL */
    :root {
        --azul-escuro: #002B49;
        --azul-medio: #005691;
        --azul-claro: #00A8E8;
        --azul-gelo: #1E293B;
        --branco: #FFFFFF;
    }

    h1, h2, h3, h4, h5 {
        color: #38BDF8 !important;
        font-weight: 700 !important;
    }

    /* TABS / ABAS */
    button[data-baseweb="tab"] {
        background-color: #0F172A !important;
        color: #94A3B8 !important;
        border-radius: 8px 8px 0 0 !important;
        padding: 10px 20px !important;
        font-weight: 600 !important;
        border: 1px solid #1E293B !important;
        margin-right: 4px !important;
    }
    button[data-baseweb="tab"][aria-selected="true"] {
        background-color: #0284C7 !important;
        color: #FFFFFF !important;
        border-bottom: 3px solid #38BDF8 !important;
    }

    /* BOTÕES */
    div.stButton > button, div.stDownloadButton > button {
        background-color: #0284C7 !important;
        color: #FFFFFF !important;
        border-radius: 6px !important;
        border: none !important;
        font-weight: 600 !important;
        transition: all 0.3s ease !important;
    }
    div.stButton > button:hover, div.stDownloadButton > button:hover {
        background-color: #38BDF8 !important;
        color: #0F172A !important;
    }

    /* KPIS METRICAS */
    div[data-testid="stMetric"] {
        background-color: #1E293B !important;
        border-left: 5px solid #0284C7 !important;
        padding: 12px 16px !important;
        border-radius: 8px !important;
    }
    div[data-testid="stMetricLabel"] {
        color: #94A3B8 !important;
        font-weight: 600 !important;
    }
    div[data-testid="stMetricValue"] {
        color: #38BDF8 !important;
        font-weight: 700 !important;
    }

    /* TABELAS ZEBRADAS DE ALTA VISIBILIDADE */
    div[data-testid="stDataFrame"] div[role="row"]:nth-child(even) {
        background-color: #1E293B !important;
    }
    div[data-testid="stDataFrame"] div[role="row"]:nth-child(odd) {
        background-color: #0F172A !important;
    }
</style>
""", unsafe_allow_html=True)

DB_NAME = "fretes_local.db"

# --- INICIALIZAÇÃO DO BANCO DE DADOS ---
def init_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS fretes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            chave_nfe TEXT UNIQUE,
            nf TEXT,
            cliente TEXT,
            cnpj_cpf TEXT,
            uf_origem TEXT,
            uf_destino TEXT,
            cidade_destino TEXT,
            valor_nf REAL,
            volume INTEGER,
            peso REAL,
            transportadora TEXT,
            cte TEXT,
            num_operacional TEXT,
            modal TEXT,
            valor_frete REAL,
            pct_frete REAL,
            natureza_orcamentaria TEXT,
            unidade TEXT,
            departamento TEXT,
            centro_custo TEXT,
            data_emissao TEXT,
            data_previsao TEXT,
            data_entrega TEXT,
            dias_previstos INTEGER,
            tipo_frete TEXT,
            aliquota_icms REAL,
            status TEXT,
            cnpj_nosso TEXT
        )
    ''')
    conn.commit()
    conn.close()

init_db()

COLUNAS_BASE = [
    'nf', 'cliente', 'cnpj_cpf', 'uf_origem', 'uf_destino', 'cidade_destino', 
    'valor_nf', 'volume', 'peso', 'transportadora', 'cte', 'num_operacional', 'modal', 
    'valor_frete', 'pct_frete', 'unidade', 'departamento', 'centro_custo', 
    'natureza_orcamentaria', 'tipo_frete', 'aliquota_icms', 'status', 
    'data_emissao', 'data_previsao', 'data_entrega', 'dias_previstos'
]

# --- FUNÇÕES AUXILIARES ---
def formatar_cnpj(cnpj):
    digits = ''.join(filter(str.isdigit, str(cnpj)))
    if len(digits) == 14:
        return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"
    elif len(digits) == 11:
        return f"{digits[:3]}.{digits[3:6]}.{digits[6:9]}-{digits[9:]}"
    return cnpj

NAT_ORCAMENTARIA = { "Aéreo": "02.02.01.02", "Rodoviário": "02.02.01.03" }

def identificar_unidade_por_cnpj(*cnpjs):
    texto_cnpjs = " ".join([str(c) for c in cnpjs if c])
    clean_cnpj = ''.join(filter(str.isdigit, texto_cnpjs))
    if "02323120000236" in clean_cnpj:
        return "Filial PB"
    elif "02323120000155" in clean_cnpj:
        return "Matriz CE"
    return "Matriz CE"

def carregar_dados():
    conn = sqlite3.connect(DB_NAME)
    df = pd.read_sql_query("SELECT * FROM fretes", conn)
    conn.close()
    if df.empty:
        return pd.DataFrame(columns=COLUNAS_BASE)
    
    if 'id' in df.columns:
        df = df.drop(columns=['id'])

    for col in ['data_emissao', 'data_previsao', 'data_entrega']:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors='coerce').dt.date

    # GARANTIR QUE DEPARTAMENTO E CENTRO_CUSTO NUNCA SEJAM NULOS
    for idx, row in df.iterrows():
        unid = row.get('unidade', '')
        if not row.get('departamento') or pd.isna(row.get('departamento')):
            df.loc[idx, 'departamento'] = "PB- HEMOTERAPIA" if "PB" in str(unid) else "CE- HEMOTERAPIA"
        if not row.get('centro_custo') or pd.isna(row.get('centro_custo')):
            df.loc[idx, 'centro_custo'] = "01.06.08" if "PB" in str(unid) else "01.01.08"

    return df

def salvar_ou_atualizar(df):
    conn = sqlite3.connect(DB_NAME)
    df_salvar = df.copy()
    for col in ['data_emissao', 'data_previsao', 'data_entrega']:
        if col in df_salvar.columns:
            df_salvar[col] = df_salvar[col].astype(str)
            
    df_salvar.to_sql('fretes', conn, if_exists='replace', index=False)
    conn.close()

# --- GERADOR DE EXCEL DO PAINEL GERAL ---
def gerar_excel_painel_geral(df_input):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Painel_Geral_Fretes"
    
    header_fill = PatternFill(start_color="002B49", end_color="002B49", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True, color="002B49")
    regular_font = Font(name="Calibri", size=11)
    
    fill_even = PatternFill(start_color="EDF4FB", end_color="EDF4FB", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    headers = [
        "Item Nº", "Cliente", "CNPJ/CPF", "NF", "Transportadora", "DACTE/CTe", 
        "LINHA", "C. CUSTO", "Modal", "Nat. Frete", 
        "Valor NF (R$)", "Valor Frete (R$)", "% Frete", "Unidade", 
        "NAT. ORÇ.", "Alíq. ICMS (%)", "Data Emissão"
    ]
    
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    current_row = 2
    item_count = 1
    for idx, row in df_input.iterrows():
        row_fill = fill_even if current_row % 2 == 0 else fill_odd
        
        ws.cell(row=current_row, column=1, value=item_count).font = bold_font
        ws.cell(row=current_row, column=2, value=str(row.get('cliente', ''))).font = regular_font
        ws.cell(row=current_row, column=3, value=str(row.get('cnpj_cpf', ''))).font = regular_font
        ws.cell(row=current_row, column=4, value=str(row.get('nf', ''))).font = regular_font
        ws.cell(row=current_row, column=5, value=str(row.get('transportadora', ''))).font = regular_font
        ws.cell(row=current_row, column=6, value=str(row.get('cte', ''))).font = regular_font
        ws.cell(row=current_row, column=7, value=str(row.get('departamento', ''))).font = regular_font
        ws.cell(row=current_row, column=8, value=str(row.get('centro_custo', ''))).font = regular_font
        ws.cell(row=current_row, column=9, value=str(row.get('modal', ''))).font = regular_font
        ws.cell(row=current_row, column=10, value=str(row.get('tipo_frete', ''))).font = regular_font
        
        c_vnf = ws.cell(row=current_row, column=11, value=float(row.get('valor_nf', 0)))
        c_vnf.number_format = 'R$ #,##0.00'
        c_vnf.font = regular_font
        
        c_vfr = ws.cell(row=current_row, column=12, value=float(row.get('valor_frete', 0)))
        c_vfr.number_format = 'R$ #,##0.00'
        c_vfr.font = regular_font
        
        c_pct = ws.cell(row=current_row, column=13, value=float(row.get('pct_frete', 0))/100.0)
        c_pct.number_format = '0.00%'
        c_pct.font = regular_font

        ws.cell(row=current_row, column=14, value=str(row.get('unidade', ''))).font = regular_font
        ws.cell(row=current_row, column=15, value=str(row.get('natureza_orcamentaria', ''))).font = regular_font
        
        c_icms = ws.cell(row=current_row, column=16, value=float(row.get('aliquota_icms', 0))/100.0)
        c_icms.number_format = '0.00%'
        c_icms.font = regular_font
        
        ws.cell(row=current_row, column=17, value=str(row.get('data_emissao', ''))).font = regular_font
        
        for c in range(1, 18):
            cell_item = ws.cell(row=current_row, column=c)
            cell_item.border = thin_border
            cell_item.fill = row_fill
            
        current_row += 1
        item_count += 1
        
    wb.save(output)
    output.seek(0)
    return output

# --- GERADOR DE EXCEL (XLSX) FINANCEIRO ---
def gerar_excel_financeiro(df_input, total_nf, total_frete, total_pct):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatorio_Financeiro"
    
    header_fill = PatternFill(start_color="002B49", end_color="002B49", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True, color="002B49")
    regular_font = Font(name="Calibri", size=11)
    
    fill_even = PatternFill(start_color="EDF4FB", end_color="EDF4FB", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    headers = [
        "Item Nº", "Transportadora", "NF", "DACTE", "LINHA", "C. CUSTO",
        "Cliente", "CNPJ/CPF", "UF Destino", "Cidade Destino",
        "Valor NF (R$)", "Valor Frete (R$)", "% Frete", "Data Emissão"
    ]
    
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    current_row = 2
    item_count = 1
    for idx, row in df_input.iterrows():
        row_fill = fill_even if current_row % 2 == 0 else fill_odd
        
        ws.cell(row=current_row, column=1, value=item_count).font = bold_font
        ws.cell(row=current_row, column=2, value=str(row.get('transportadora', ''))).font = regular_font
        ws.cell(row=current_row, column=3, value=str(row.get('nf', ''))).font = regular_font
        ws.cell(row=current_row, column=4, value=str(row.get('cte', ''))).font = regular_font
        ws.cell(row=current_row, column=5, value=str(row.get('departamento', ''))).font = regular_font
        ws.cell(row=current_row, column=6, value=str(row.get('centro_custo', ''))).font = regular_font
        ws.cell(row=current_row, column=7, value=str(row.get('cliente', ''))).font = regular_font
        ws.cell(row=current_row, column=8, value=str(row.get('cnpj_cpf', ''))).font = regular_font
        ws.cell(row=current_row, column=9, value=str(row.get('uf_destino', ''))).font = regular_font
        ws.cell(row=current_row, column=10, value=str(row.get('cidade_destino', ''))).font = regular_font
        
        c_vnf = ws.cell(row=current_row, column=11, value=float(row.get('valor_nf', 0)))
        c_vnf.number_format = 'R$ #,##0.00'
        c_vnf.font = regular_font
        
        c_vfr = ws.cell(row=current_row, column=12, value=float(row.get('valor_frete', 0)))
        c_vfr.number_format = 'R$ #,##0.00'
        c_vfr.font = regular_font
        
        c_pct = ws.cell(row=current_row, column=13, value=float(row.get('pct_frete', 0))/100.0)
        c_pct.number_format = '0.00%'
        c_pct.font = regular_font
        
        ws.cell(row=current_row, column=14, value=str(row.get('data_emissao', ''))).font = regular_font
        
        for c in range(1, 15):
            cell_item = ws.cell(row=current_row, column=c)
            cell_item.border = thin_border
            cell_item.fill = row_fill
            
        current_row += 1
        item_count += 1
        
    ws.cell(row=current_row, column=1, value="TOTAL").font = bold_font
    c_tot_vnf = ws.cell(row=current_row, column=11, value=float(total_nf))
    c_tot_vnf.number_format = 'R$ #,##0.00'
    c_tot_vnf.font = bold_font
    
    c_tot_vfr = ws.cell(row=current_row, column=12, value=float(total_frete))
    c_tot_vfr.number_format = 'R$ #,##0.00'
    c_tot_vfr.font = bold_font
    
    c_tot_pct = ws.cell(row=current_row, column=13, value=float(total_pct)/100.0)
    c_tot_pct.number_format = '0.00%'
    c_tot_pct.font = bold_font
    
    for c in range(1, 15):
        ws.cell(row=current_row, column=c).border = thin_border

    wb.save(output)
    output.seek(0)
    return output

# --- GERADOR DE EXCEL (XLSX) PAGAMENTOS ---
def gerar_excel_pagamento(df_input, total_frete, fatura_info):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fatura_Financeiro"
    
    header_fill = PatternFill(start_color="002B49", end_color="002B49", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True, color="002B49")
    regular_font = Font(name="Calibri", size=11)
    
    fill_even = PatternFill(start_color="EDF4FB", end_color="EDF4FB", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    ws.cell(row=1, column=1, value="Transp.").font = bold_font
    ws.cell(row=1, column=2, value=str(fatura_info.get("transp", "")))
    
    ws.cell(row=2, column=1, value="Fatura").font = bold_font
    ws.cell(row=2, column=2, value=str(fatura_info.get("num_fatura", "")))
    
    ws.cell(row=3, column=1, value="Venc.").font = bold_font
    ws.cell(row=3, column=2, value=str(fatura_info.get("venc", "")))
    
    ws.cell(row=4, column=1, value="Valor Total").font = bold_font
    c_tot = ws.cell(row=4, column=2, value=float(total_frete))
    c_tot.number_format = 'R$ #,##0.00'
    c_tot.font = bold_font

    headers = [
        "Item Nº", "Transportadora", "NF", "LINHA", 
        "DACTE", "NAT. ORÇ.", "C. CUSTO", "NAT FRETE", "Alíq. ICMS (%)", "Valor Frete (R$)"
    ]
    
    start_row = 6
    for col_num, h_text in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    current_row = start_row + 1
    item_count = 1
    for _, row in df_input.iterrows():
        row_fill = fill_even if current_row % 2 == 0 else fill_odd
        
        ws.cell(row=current_row, column=1, value=item_count).font = bold_font
        ws.cell(row=current_row, column=2, value=str(row.get('transportadora', ''))).font = regular_font
        ws.cell(row=current_row, column=3, value=str(row.get('nf', ''))).font = regular_font
        ws.cell(row=current_row, column=4, value=str(row.get('departamento', ''))).font = regular_font
        ws.cell(row=current_row, column=5, value=str(row.get('cte', ''))).font = regular_font
        ws.cell(row=current_row, column=6, value=str(row.get('natureza_orcamentaria', ''))).font = regular_font
        ws.cell(row=current_row, column=7, value=str(row.get('centro_custo', ''))).font = regular_font
        ws.cell(row=current_row, column=8, value=str(row.get('tipo_frete', ''))).font = regular_font
        
        c_icms = ws.cell(row=current_row, column=9, value=float(row.get('aliquota_icms', 0))/100.0)
        c_icms.number_format = '0.00%'
        c_icms.font = regular_font
        
        c_vfr = ws.cell(row=current_row, column=10, value=float(row.get('valor_frete', 0)))
        c_vfr.number_format = 'R$ #,##0.00'
        c_vfr.font = regular_font
        
        for c in range(1, 11):
            cell_item = ws.cell(row=current_row, column=c)
            cell_item.border = thin_border
            cell_item.fill = row_fill
            
        current_row += 1
        item_count += 1
        
    ws.cell(row=current_row, column=1, value="TOTAL").font = bold_font
    c_tot_vfr = ws.cell(row=current_row, column=10, value=float(total_frete))
    c_tot_vfr.number_format = 'R$ #,##0.00'
    c_tot_vfr.font = bold_font
    
    for c in range(1, 11):
        ws.cell(row=current_row, column=c).border = thin_border

    wb.save(output)
    output.seek(0)
    return output

# --- GERADOR DE PDF DA FATURA ---
def gerar_pdf_fatura(transp_nome, num_fatura, venc_str, df_selecionados, total_fatura):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
    story = []
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=15, textColor=colors.HexColor('#002B49'), spaceAfter=8)
    
    cell_style = ParagraphStyle('CellText', parent=styles['Normal'], fontSize=7, leading=8, alignment=1)
    cell_head = ParagraphStyle('HeadText', parent=styles['Normal'], fontSize=7.5, leading=8.5, alignment=1, textColor=colors.white, fontName='Helvetica-Bold')

    story.append(Paragraph("<b>RELATÓRIO DE FATURA PARA PAGAMENTO DE FRETE</b>", title_style))
    story.append(Spacer(1, 6))
    
    val_tot_str = f"R$ {total_fatura:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    cabecalho_data = [
        [Paragraph(f"<b>Transportadora:</b> {transp_nome}", styles['Normal']), Paragraph(f"<b>Nº Fatura:</b> {num_fatura}", styles['Normal'])],
        [Paragraph(f"<b>Vencimento:</b> {venc_str}", styles['Normal']), Paragraph(f"<b>Valor Total:</b> {val_tot_str}", styles['Normal'])]
    ]
    t_cab = Table(cabecalho_data, colWidths=[370, 370])
    t_cab.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EDF4FB')),
        ('PADDING', (0,0), (-1,-1), 6),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1'))
    ]))
    story.append(t_cab)
    story.append(Spacer(1, 10))
    
    cols = ['Nº', 'Transportadora', 'NF', 'LINHA', 'DACTE', 'NAT. ORÇ.', 'C. CUSTO', 'NAT FRETE', 'Alíq. ICMS', 'Valor Frete (R$)']
    table_data = [[Paragraph(c, cell_head) for c in cols]]
    
    item_idx = 1
    for idx, row in df_selecionados.iterrows():
        val_fmt = f"R$ {row['valor_frete']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        aliq_fmt = f"{row['aliquota_icms']:.2f}%"
        
        table_data.append([
            Paragraph(str(item_idx), cell_style),
            Paragraph(str(row['transportadora']), cell_style),
            Paragraph(str(row['nf']), cell_style),
            Paragraph(str(row['departamento']), cell_style),
            Paragraph(str(row['cte']), cell_style),
            Paragraph(str(row['natureza_orcamentaria']), cell_style),
            Paragraph(str(row['centro_custo']), cell_style),
            Paragraph(str(row['tipo_frete']), cell_style),
            Paragraph(aliq_fmt, cell_style),
            Paragraph(val_fmt, cell_style)
        ])
        item_idx += 1
        
    table_data.append([
        Paragraph("<b>TOTAL</b>", cell_style),
        "", "", "", "", "", "", "", "",
        Paragraph(f"<b>{val_tot_str}</b>", cell_style)
    ])
    
    t_dados = Table(table_data, colWidths=[30, 160, 55, 100, 60, 65, 55, 55, 55, 105])
    t_dados.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#002B49')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#EDF4FB')]),
        ('SPAN', (0, -1), (8, -1)),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#D0E2F3'))
    ]))
    story.append(t_dados)
    
    doc.build(story)
    buffer.seek(0)
    return buffer

# --- GERADOR DE PDF DO FINANCEIRO ---
def gerar_pdf_financeiro(df_input, total_nf, total_frete, total_pct):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
    story = []
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=15, textColor=colors.HexColor('#002B49'), spaceAfter=8)
    
    cell_style = ParagraphStyle('CellText', parent=styles['Normal'], fontSize=7, leading=8, alignment=1)
    cell_head = ParagraphStyle('HeadText', parent=styles['Normal'], fontSize=7.5, leading=8.5, alignment=1, textColor=colors.white, fontName='Helvetica-Bold')

    story.append(Paragraph("<b>RELATÓRIO GERENCIAL FINANCEIRO DE FRETES</b>", title_style))
    story.append(Spacer(1, 6))
    
    tot_nf_str = f"R$ {total_nf:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    tot_fr_str = f"R$ {total_frete:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    pct_str = f"{total_pct:.2f}%"
    
    cabecalho_data = [
        [Paragraph(f"<b>Valor Total NFs:</b> {tot_nf_str}", styles['Normal']), Paragraph(f"<b>Valor Total Fretes:</b> {tot_fr_str}", styles['Normal']), Paragraph(f"<b>% Representativa:</b> {pct_str}", styles['Normal'])]
    ]
    t_cab = Table(cabecalho_data, colWidths=[240, 240, 260])
    t_cab.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#EDF4FB')),
        ('PADDING', (0,0), (-1,-1), 6),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1'))
    ]))
    story.append(t_cab)
    story.append(Spacer(1, 10))
    
    cols = ['Nº', 'Transportadora', 'NF', 'DACTE', 'LINHA', 'C. CUSTO', 'Cliente', 'Valor NF', 'Valor Frete', '% Frete']
    table_data = [[Paragraph(c, cell_head) for c in cols]]
    
    item_idx = 1
    for idx, row in df_input.iterrows():
        vnf_fmt = f"R$ {row['valor_nf']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        vfr_fmt = f"R$ {row['valor_frete']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        pct_fmt = f"{row['pct_frete']:.2f}%"
        table_data.append([
            Paragraph(str(item_idx), cell_style),
            Paragraph(str(row['transportadora']), cell_style),
            Paragraph(str(row['nf']), cell_style),
            Paragraph(str(row['cte']), cell_style),
            Paragraph(str(row['departamento']), cell_style),
            Paragraph(str(row['centro_custo']), cell_style),
            Paragraph(str(row['cliente']), cell_style),
            Paragraph(vnf_fmt, cell_style),
            Paragraph(vfr_fmt, cell_style),
            Paragraph(pct_fmt, cell_style)
        ])
        item_idx += 1
        
    table_data.append([
        Paragraph("<b>TOTAL</b>", cell_style),
        "", "", "", "", "", "",
        Paragraph(f"<b>{tot_nf_str}</b>", cell_style),
        Paragraph(f"<b>{tot_fr_str}</b>", cell_style),
        Paragraph(f"<b>{pct_str}</b>", cell_style)
    ])
    
    t_dados = Table(table_data, colWidths=[30, 130, 45, 55, 90, 50, 130, 80, 80, 50])
    t_dados.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#002B49')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#EDF4FB')]),
        ('SPAN', (0, -1), (6, -1)),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#D0E2F3'))
    ]))
    story.append(t_dados)
    
    doc.build(story)
    buffer.seek(0)
    return buffer

# --- PARSER DE CT-E COM RELEITURA DE MODAL MISTO ---
def parse_xml_cte_exclusivo(source):
    try:
        tree = ET.parse(source)
        root = tree.getroot()
        ns_cte = {'cte': 'http://www.portalfiscal.inf.br/cte'}
        
        infCte = root.find('.//cte:infCte', ns_cte)
        if infCte is None:
            return None, "Arquivo não é um XML válido de CT-e."

        n_cte = root.find('.//cte:ide/cte:nCT', ns_cte) or root.find('.//cte:nCT', ns_cte)
        v_rec = root.find('.//cte:vPrest/cte:vTPrest', ns_cte)
        aliq_icms = root.find('.//cte:imp/cte:ICMS//cte:pICMS', ns_cte)
        
        modal_code = root.find('.//cte:ide/cte:modal', ns_cte)
        code_val = modal_code.text if modal_code is not None else ''
        
        obs_textos = [obs.text.upper() for obs in root.findall('.//cte:compl/cte:ObsCont/cte:xTexto', ns_cte) if obs.text]
        obs_geral = " ".join(obs_textos)
        
        has_aereo_group = root.find('.//cte:infCTeNorm/cte:infModal/cte:aereo', ns_cte) is not None

        if code_val == "02" or has_aereo_group or "AEREO" in obs_geral or "AÉREO" in obs_geral:
            modal_str = "Aéreo"
        else:
            modal_str = "Rodoviário"

        n_minu = root.find('.//cte:aereo/cte:nMinu', ns_cte) or root.find('.//cte:aereo/cte:nOCA', ns_cte)
        num_operacional = n_minu.text if n_minu is not None else ''

        d_prev_aereo = root.find('.//cte:dPrevAereo', ns_cte)
        d_prog_rodo = root.find('.//cte:compl/cte:Entrega/cte:comData/cte:dProg', ns_cte)
        d_emi = root.find('.//cte:dhEmi', ns_cte)
        
        d_prev_str = ''
        if d_prev_aereo is not None and d_prev_aereo.text:
            d_prev_str = d_prev_aereo.text
        elif d_prog_rodo is not None and d_prog_rodo.text:
            d_prev_str = d_prog_rodo.text
        elif d_emi is not None and d_emi.text:
            d_prev_str = d_emi.text[:10]
        else:
            d_prev_str = str(datetime.now().date())

        transp = root.find('.//cte:emit/cte:xNome', ns_cte)
        dest = root.find('.//cte:dest/cte:xNome', ns_cte)
        
        dest_cnpj = root.find('.//cte:dest/cte:CNPJ', ns_cte)
        dest_cpf = root.find('.//cte:dest/cte:CPF', ns_cte)
        toma_cnpj = root.find('.//cte:toma4/cte:CNPJ', ns_cte)
        rem_cnpj = root.find('.//cte:rem/cte:CNPJ', ns_cte)
        
        raw_cnpj_dest = dest_cnpj.text if (dest_cnpj is not None and dest_cnpj.text) else ''
        raw_toma_cnpj = toma_cnpj.text if (toma_cnpj is not None and toma_cnpj.text) else ''
        raw_rem_cnpj = rem_cnpj.text if (rem_cnpj is not None and rem_cnpj.text) else ''

        cnpj_cliente_val = formatar_cnpj(raw_cnpj_dest if raw_cnpj_dest else raw_toma_cnpj)

        unidade_auto = identificar_unidade_por_cnpj(raw_cnpj_dest, raw_toma_cnpj, raw_rem_cnpj)
        
        clean_dest = ''.join(filter(str.isdigit, raw_cnpj_dest + raw_toma_cnpj))
        if "02323120000155" in clean_dest or "02323120000236" in clean_dest:
            tipo_frete_val = "FOB"
        else:
            tipo_frete_val = "CIF"

        uf_dest = root.find('.//cte:UFFim', ns_cte)
        cidade_dest = root.find('.//cte:xMunFim', ns_cte)
        uf_orig = root.find('.//cte:UFIni', ns_cte)

        v_carga = root.find('.//cte:vCarga', ns_cte)
        val_carga_num = float(v_carga.text) if (v_carga is not None and v_carga.text) else 0.0

        q_vol = 0
        q_peso = 0.0
        for infQ in root.findall('.//cte:infCarga/cte:infQ', ns_cte):
            tpMed = infQ.find('.//cte:tpMed', ns_cte)
            qCarga = infQ.find('.//cte:qCarga', ns_cte)
            if tpMed is not None and qCarga is not None and qCarga.text:
                med_upper = tpMed.text.upper()
                if 'VOL' in med_upper or 'CAIXA' in med_upper:
                    q_vol = int(float(qCarga.text))
                elif 'PESO' in med_upper or 'KG' in med_upper:
                    q_peso = float(qCarga.text)

        nfs_vinculadas = []
        for infNFe in root.findall('.//cte:infDoc/cte:infNFe', ns_cte):
            chave_elem = infNFe.find('.//cte:chave', ns_cte)
            if chave_elem is not None and chave_elem.text:
                chave = chave_elem.text
                num_nf = str(int(chave[25:34])) if len(chave) >= 34 else '0'
                nfs_vinculadas.append((chave, num_nf))

        string_nfs = "/".join([n[1] for n in nfs_vinculadas]) if nfs_vinculadas else '0'
        chave_principal = nfs_vinculadas[0][0] if nfs_vinculadas else f"CTE_{n_cte.text if n_cte is not None else '0'}"

        v_frete_val = float(v_rec.text) if (v_rec is not None and v_rec.text) else 0.0
        pct_f = round((v_frete_val / val_carga_num) * 100, 2) if val_carga_num > 0 else 0.0

        depto_padrao = "PB- HEMOTERAPIA" if "Filial" in unidade_auto else "CE- HEMOTERAPIA"
        cc_padrao = "01.06.08" if "Filial" in unidade_auto else "01.01.08"

        return {
            'chave_nfe': chave_principal,
            'nf': string_nfs,
            'cliente': dest.text if dest is not None else '',
            'cnpj_cpf': cnpj_cliente_val,
            'uf_origem': uf_orig.text if uf_orig is not None else '',
            'uf_destino': uf_dest.text if uf_dest is not None else '',
            'cidade_destino': cidade_dest.text if cidade_dest is not None else '',
            'valor_nf': val_carga_num,
            'volume': q_vol,
            'peso': q_peso,
            'transportadora': transp.text if transp is not None else 'NÃO INFORMADA',
            'cte': n_cte.text if n_cte is not None else '',
            'num_operacional': num_operacional,
            'modal': modal_str,
            'valor_frete': v_frete_val,
            'pct_frete': pct_f,
            'natureza_orcamentaria': NAT_ORCAMENTARIA.get(modal_str, '02.02.01.03'),
            'unidade': unidade_auto,
            'departamento': depto_padrao,
            'centro_custo': cc_padrao,
            'data_emissao': d_emi.text[:10] if d_emi is not None else str(datetime.now().date()),
            'data_previsao': d_prev_str,
            'data_entrega': None,
            'dias_previstos': 0,
            'tipo_frete': tipo_frete_val,
            'aliquota_icms': float(aliq_icms.text) if (aliq_icms is not None and aliq_icms.text) else 0.0,
            'status': 'PROCESSADO VIA CTE',
            'cnpj_nosso': raw_rem_cnpj
        }, None
    except Exception as e:
        return None, f"Erro ao processar CT-e: {e}"

# --- INTERFACE PRINCIPAL ---
st.title("📦 Central de Rastreabilidade, Fretes & Faturas")

aba_import, aba_gestao, aba_rel_fin, aba_rel_pag, aba_amostragem = st.tabs([
    "📥 Importar CT-es",
    "📋 Painel de Gestão",
    "📊 Relatório Financeiro",
    "💳 Pagamentos",
    "📈 Amostragem"
])

# --- ABA 1: IMPORTAÇÃO ---
with aba_import:
    st.header("Importação Direta de CT-es")
    st.info("Carregue os arquivos XML dos CT-es. O sistema identifica automaticamente Matriz/Filial, CIF/FOB, pesos, valores e CNPJs.")

    # INICIALIZAÇÃO DA CHAVE DE RESET DOS UPLOADS DE ARQUIVOS
    if 'uploader_key' not in st.session_state:
        st.session_state['uploader_key'] = 0

    files = st.file_uploader(
        "Arraste ou selecione os arquivos XML dos CT-es", 
        type=["xml"], 
        accept_multiple_files=True,
        key=f"uploader_{st.session_state['uploader_key']}"
    )

    c_proc, c_limp = st.columns([2, 2])

    with c_proc:
        if files and st.button("🚀 Processar CT-es Carregados"):
            df_base = carregar_dados()
            processados = 0

            for f in files:
                dados, err = parse_xml_cte_exclusivo(f)
                if err:
                    st.error(f"Arquivo '{f.name}': {err}")
                elif dados:
                    if not df_base.empty and dados['chave_nfe'] in df_base['chave_nfe'].values:
                        idx = df_base[df_base['chave_nfe'] == dados['chave_nfe']].index[0]
                        for k, v in dados.items():
                            df_base.loc[idx, k] = v
                        salvar_ou_atualizar(df_base)
                    else:
                        conn = sqlite3.connect(DB_NAME)
                        pd.DataFrame([dados]).to_sql('fretes', conn, if_exists='append', index=False)
                        conn.close()
                        df_base = carregar_dados()
                    processados += 1

            st.success(f"Sucesso! {processados} CT-e(s) importados e registrados no banco de dados.")

    # BOTÃO PARA EXCLUIR/LIMPAR APENAS OS ARQUIVOS ANEXADOS NA TELA
    with c_limp:
        if files:
            if st.button("🗑️ Limpar Seleção / Remover Arquivos da Tela"):
                st.session_state['uploader_key'] += 1
                st.rerun()

# --- ABA 2: PAINEL DE GESTÃO ---
with aba_gestao:
    st.header("Painel de Gestão Geral")
    df = carregar_dados()

    if not df.empty:
        busca = st.text_input("🔍 Pesquisar em qualquer coluna (digite número da NF, cliente, CTE, etc.):", "")
        if busca:
            mask = df.astype(str).apply(lambda row: row.str.contains(busca, case=False).any(), axis=1)
            df = df[mask]

        for idx, row in df.iterrows():
            df.loc[idx, 'unidade'] = identificar_unidade_por_cnpj(row.get('cnpj_nosso', ''), row.get('cnpj_cpf', ''))
            df.loc[idx, 'natureza_orcamentaria'] = NAT_ORCAMENTARIA.get(row['modal'], "02.02.01.03")
            
            if pd.notnull(row['data_previsao']) and pd.notnull(row['data_entrega']):
                d1 = pd.to_datetime(row['data_previsao'])
                d2 = pd.to_datetime(row['data_entrega'])
                df.loc[idx, 'dias_previstos'] = int((d2 - d1).days)
            else:
                df.loc[idx, 'dias_previstos'] = 0

            if row['valor_nf'] > 0 and row['valor_frete'] > 0:
                df.loc[idx, 'pct_frete'] = round((row['valor_frete'] / row['valor_nf']) * 100, 2)

    colunas_ordem_painel = [
        'cliente', 'cnpj_cpf', 'nf', 'transportadora', 'cte', 
        'departamento', 'centro_custo', 'modal', 'tipo_frete', 'valor_nf', 
        'valor_frete', 'pct_frete', 'unidade', 'natureza_orcamentaria', 
        'aliquota_icms', 'data_emissao', 'data_previsao', 'data_entrega', 'dias_previstos'
    ]
    cols_existentes = [c for c in colunas_ordem_painel if c in df.columns]

    df_editado = st.data_editor(
        df[cols_existentes],
        column_config={
            "cliente": st.column_config.TextColumn("Cliente / Fornecedor", disabled=True),
            "cnpj_cpf": st.column_config.TextColumn("CNPJ/CPF Cliente", disabled=True),
            "nf": st.column_config.TextColumn("NF(s)", disabled=True),
            "transportadora": st.column_config.TextColumn("Transportadora", disabled=True),
            "cte": st.column_config.TextColumn("DACTE / CTe", disabled=True),
            "departamento": st.column_config.SelectboxColumn(
                "LINHA (Departamento)", 
                options=[
                    "PB- HEMOTERAPIA", "PB- HOSPITALAR", "PB- ORSENSE",
                    "CE- HEMOTERAPIA", "CE- HOSPITALAR", "CE- ORSENSE"
                ],
                required=True
            ),
            "centro_custo": st.column_config.SelectboxColumn(
                "C. CUSTO (Edição Manual)",
                options=[
                    "01.06.08", "01.06.09", "01.06.10",
                    "01.01.08", "01.01.09", "01.01.10"
                ],
                required=True
            ),
            "modal": st.column_config.TextColumn("Modal", disabled=True),
            "tipo_frete": st.column_config.TextColumn("Nat. Frete", disabled=True),
            "valor_nf": st.column_config.NumberColumn("Valor NF", format="R$ %.2f", disabled=True),
            "valor_frete": st.column_config.NumberColumn("Valor Frete", format="R$ %.2f", disabled=True),
            "pct_frete": st.column_config.NumberColumn("% Frete", format="%.2f %%", disabled=True),
            "unidade": st.column_config.TextColumn("Unidade", disabled=True),
            "natureza_orcamentaria": st.column_config.TextColumn("NAT. ORÇ. (Auto)", disabled=True),
            "aliquota_icms": st.column_config.NumberColumn("Alíq. ICMS", format="%.2f %%", disabled=True),
            "data_emissao": st.column_config.DateColumn("Data Emissão", disabled=True),
            "data_previsao": st.column_config.DateColumn("Previsão Entrega", disabled=True),
            "data_entrega": st.column_config.DateColumn("Data de Entrega"),
            "dias_previstos": st.column_config.NumberColumn("Diferença Dias", format="%d dia(s)", disabled=True)
        },
        use_container_width=True,
        num_rows="fixed",
        height=550
    )

    c_b1, c_b2 = st.columns(2)
    with c_b1:
        if st.button("💾 Salvar Alterações de Linha e Centro de Custo"):
            df_full_original = carregar_dados()
            df_salvar_limpo = df_editado.copy()
            for idx_s, r_s in df_salvar_limpo.iterrows():
                if not r_s.get('departamento') or pd.isna(r_s.get('departamento')):
                    df_salvar_limpo.loc[idx_s, 'departamento'] = "PB- HEMOTERAPIA" if "PB" in str(r_s.get('unidade')) else "CE- HEMOTERAPIA"
                if not r_s.get('centro_custo') or pd.isna(r_s.get('centro_custo')):
                    df_salvar_limpo.loc[idx_s, 'centro_custo'] = "01.06.08" if "PB" in str(r_s.get('unidade')) else "01.01.08"

            for col_f in df_full_original.columns:
                if col_f not in df_salvar_limpo.columns:
                    df_salvar_limpo[col_f] = df_full_original[col_f]

            salvar_ou_atualizar(df_salvar_limpo)
            st.success("Alterações salvas com sucesso!")

    # EXTRATOR DE RELATÓRIO DO PAINEL GERAL
    with c_b2:
        if not df_editado.empty:
            excel_painel_bytes = gerar_excel_painel_geral(df_editado)
            st.download_button(
                label="📊 Exportar Painel Geral Completo em Excel (.xlsx)",
                data=excel_painel_bytes,
                file_name="Painel_Geral_Fretes_Completo.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

# --- ABA 3: RELATÓRIO FINANCEIRO ---
with aba_rel_fin:
    st.header("Relatório Financeiro de Fretes")
    df = carregar_dados()

    st.subheader("Filtros do Relatório")
    col_r1, col_r2, col_r3, col_r4 = st.columns(4)
    
    transp_opts = list(df['transportadora'].dropna().unique()) if not df.empty else []
    cli_opts = list(df['cliente'].dropna().unique()) if not df.empty else []

    with col_r1:
        f_inicio = st.date_input("Data Inicial (Emissão)", value=None, key="rf_ini")
    with col_r2:
        f_fim = st.date_input("Data Final (Emissão)", value=None, key="rf_fim")
    with col_r3:
        transp_sel = st.multiselect("Transportadora", options=transp_opts)
    with col_r4:
        cli_sel = st.multiselect("Cliente", options=cli_opts)

    df_filtrado = df.copy()
    if f_inicio and not df_filtrado.empty:
        df_filtrado = df_filtrado[df_filtrado['data_emissao'] >= f_inicio]
    if f_fim and not df_filtrado.empty:
        df_filtrado = df_filtrado[df_filtrado['data_emissao'] <= f_fim]
    if transp_sel and not df_filtrado.empty:
        df_filtrado = df_filtrado[df_filtrado['transportadora'].isin(transp_sel)]
    if cli_sel and not df_filtrado.empty:
        df_filtrado = df_filtrado[df_filtrado['cliente'].isin(cli_sel)]

    tot_v_nf = df_filtrado['valor_nf'].sum() if 'valor_nf' in df_filtrado.columns else 0.0
    tot_v_frete = df_filtrado['valor_frete'].sum() if 'valor_frete' in df_filtrado.columns else 0.0
    pct_total = (tot_v_frete / tot_v_nf * 100) if tot_v_nf > 0 else 0.0

    m1, m2, m3 = st.columns(3)
    m1.metric("Valor Total NFs", f"R$ {tot_v_nf:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
    m2.metric("Valor Total Fretes", f"R$ {tot_v_frete:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
    m3.metric("% Total Representativa", f"{pct_total:.2f}%")

    cols_fin = [c for c in ['transportadora', 'nf', 'cte', 'departamento', 'centro_custo', 'cliente', 'cnpj_cpf', 'uf_destino', 'cidade_destino', 'valor_nf', 'valor_frete', 'pct_frete', 'data_emissao'] if c in df_filtrado.columns]
    
    st.dataframe(df_filtrado[cols_fin], use_container_width=True, height=500)

    if not df_filtrado.empty:
        c_fin_exp1, c_fin_exp2 = st.columns(2)
        
        excel_fin_bytes = gerar_excel_financeiro(df_filtrado[cols_fin], tot_v_nf, tot_v_frete, pct_total)
        with c_fin_exp1:
            st.download_button(
                label="📊 Exportar Relatório Financeiro em Excel (.xlsx)",
                data=excel_fin_bytes,
                file_name="Relatorio_Financeiro_Fretes.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        if HAS_REPORTLAB:
            pdf_fin_bytes = gerar_pdf_financeiro(df_filtrado[cols_fin], tot_v_nf, tot_v_frete, pct_total)
            with c_fin_exp2:
                st.download_button(
                    label="📄 Exportar Relatório Financeiro em PDF",
                    data=pdf_fin_bytes,
                    file_name="Relatorio_Financeiro_Fretes.pdf",
                    mime="application/pdf"
                )

# --- ABA 4: PAGAMENTOS ---
with aba_rel_pag:
    st.header("Envio de Fatura para Pagamento ao Financeiro")
    df = carregar_dados()

    st.subheader("1. Filtros para Busca dos CT-es")
    col_p1, col_p2, col_p3 = st.columns(3)
    transp_list = ["Todas"] + (list(df['transportadora'].dropna().unique()) if not df.empty else [])
    
    with col_p1:
        f_transp = st.selectbox("Selecione a Transportadora", options=transp_list)
    with col_p2:
        d_inicio = st.date_input("Data Emissão (Início)", value=None, key="pag_ini")
    with col_p3:
        d_fim = st.date_input("Data Emissão (Fim)", value=None, key="pag_fim")

    df_lote = df.copy()
    if f_transp != "Todas" and not df_lote.empty:
        df_lote = df_lote[df_lote['transportadora'] == f_transp]
    
    if d_inicio and not df_lote.empty:
        df_lote = df_lote[df_lote['data_emissao'] >= d_inicio]
    if d_fim and not df_lote.empty:
        df_lote = df_lote[df_lote['data_emissao'] <= d_fim]

    st.subheader("2. Tabela de Seleção dos CT-es para a Fatura")
    df_lote['SELECIONAR'] = False

    cols_pag = [c for c in ['SELECIONAR', 'transportadora', 'nf', 'departamento', 'cte', 'natureza_orcamentaria', 'centro_custo', 'tipo_frete', 'aliquota_icms', 'valor_frete'] if c in df_lote.columns]
    
    df_edit_lote = st.data_editor(
        df_lote[cols_pag],
        column_config={
            "SELECIONAR": st.column_config.CheckboxColumn("SELECIONAR"),
            "transportadora": "Transportadora",
            "nf": "NF",
            "departamento": "LINHA",
            "cte": "DACTE",
            "natureza_orcamentaria": "NAT. ORÇ.",
            "centro_custo": "C. CUSTO",
            "tipo_frete": "NAT FRETE",
            "aliquota_icms": st.column_config.NumberColumn("Alíq. ICMS (%)", format="%.2f %%"),
            "valor_frete": st.column_config.NumberColumn("Valor Frete (R$)", format="R$ %.2f")
        },
        use_container_width=True,
        height=380
    )

    selecionados = df_edit_lote[df_edit_lote['SELECIONAR'] == True] if 'SELECIONAR' in df_edit_lote.columns else pd.DataFrame()

    if not selecionados.empty:
        st.markdown("---")
        st.subheader("3. Preenchimento de Dados da Fatura")
        
        c_f1, c_f2 = st.columns(2)
        with c_f1:
            num_fatura_manual = st.text_input("Nº da Fatura", value="0474236/2026")
        with c_f2:
            venc_fatura_manual = st.date_input("Data de Vencimento da Fatura", value=datetime.now().date())

        total_fatura = selecionados['valor_frete'].sum()
        transp_nome = f_transp if f_transp != "Todas" else (selecionados['transportadora'].iloc[0] if 'transportadora' in selecionados.columns else "TRANSPORTADORA")
        venc_str = venc_fatura_manual.strftime("%d/%m/%Y") if venc_fatura_manual else "A DEFINIR"
        val_fmt_brl = f"R$ {total_fatura:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

        st.markdown("### 📄 Resumo da Fatura")
        col_m1, col_m2, col_m3, col_m4 = st.columns(4)
        col_m1.metric("Transportadora", transp_nome)
        col_m2.metric("Nº da Fatura", num_fatura_manual)
        col_m3.metric("Vencimento", venc_str)
        col_m4.metric("Valor Total da Fatura", val_fmt_brl)

        c_exp1, c_exp2 = st.columns(2)

        fatura_info_dict = {
            "transp": transp_nome,
            "num_fatura": num_fatura_manual,
            "venc": venc_str
        }
        excel_pag_bytes = gerar_excel_pagamento(selecionados, total_fatura, fatura_info_dict)

        with c_exp1:
            st.download_button(
                label="📊 Exportar Fatura em Excel (.xlsx)",
                data=excel_pag_bytes,
                file_name=f"Fatura_Financeiro_{transp_nome}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        if HAS_REPORTLAB:
            pdf_buffer = gerar_pdf_fatura(transp_nome, num_fatura_manual, venc_str, selecionados, total_fatura)
            with c_exp2:
                st.download_button(
                    label="📄 Exportar Fatura em PDF",
                    data=pdf_buffer,
                    file_name=f"Fatura_Financeiro_{transp_nome}.pdf",
                    mime="application/pdf"
                )
    else:
        st.warning("⚠️ Marque pelo menos 1 CT-e na caixa 'SELECIONAR' acima para preencher a fatura e habilitar os downloads.")

# --- ABA 5: AMOSTRAGEM / DASHBOARD EXECUTIVO ---
with aba_amostragem:
    st.header("📈 Amostragem & Resumo Operacional Geral")
    st.info("Esta aba consolida a amostragem lendo todos os CT-es cadastrados no Painel de Gestão.")

    df_amostra = carregar_dados()

    if not df_amostra.empty:
        tot_faturado_geral = df_amostra['valor_nf'].sum()
        tot_frete_geral = df_amostra['valor_frete'].sum()

        m1_am, m2_am = st.columns(2)
        m1_am.metric("Valor Faturado Total (NFs)", f"R$ {tot_faturado_geral:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
        m2_am.metric("Valor Frete Total", f"R$ {tot_frete_geral:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))

        st.markdown("---")
        st.subheader("🍩 Distribuição por Modal e Volume de NFs Enviadas")

        c_g1, c_g2 = st.columns(2)

        cores_azul = ['#003865', '#007ACC', '#00A8E8', '#4299E1', '#63B3ED', '#90CDF4']

        # 1. GRÁFICO DE ROSCA: MODAL
        with c_g1:
            st.markdown("##### 1. Distribuição de Frete por Modal")
            df_modal = df_amostra.groupby('modal')['valor_frete'].sum().reset_index()
            
            fig_m, ax_m = plt.subplots(figsize=(5, 3.8))
            wedges, texts, autotexts = ax_m.pie(
                df_modal['valor_frete'], 
                autopct='%1.1f%%', 
                startangle=90, 
                colors=cores_azul[:len(df_modal)],
                wedgeprops=dict(width=0.4, edgecolor='w'),
                pctdistance=0.75
            )
            plt.setp(autotexts, size=10, weight="bold", color="white")
            
            ax_m.legend(wedges, df_modal['modal'], title="Modal", loc="upper center", bbox_to_anchor=(0.5, -0.05), ncol=2, fontsize=8)
            ax_m.axis('equal')
            plt.tight_layout()
            st.pyplot(fig_m)

            df_modal_view = df_modal.copy()
            df_modal_view['valor_frete'] = df_modal_view['valor_frete'].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
            st.dataframe(df_modal_view.rename(columns={'modal': 'Modal', 'valor_frete': 'Valor em Frete (R$)'}), use_container_width=True)

        # 2. GRÁFICO DE ROSCA: QTD DE VIAGENS / NFS POR TRANSPORTADORA
        with c_g2:
            st.markdown("##### 2. Quantidade de Viagens / NFs por Transportadora")
            df_transp_nf = df_amostra.groupby('transportadora')['nf'].count().reset_index().rename(columns={'nf': 'qtd_nfs'})
            
            fig_t, ax_t = plt.subplots(figsize=(5, 3.8))
            wedges_t, texts_t, autotexts_t = ax_t.pie(
                df_transp_nf['qtd_nfs'], 
                autopct='%1.1f%%', 
                startangle=140, 
                colors=cores_azul[:len(df_transp_nf)],
                wedgeprops=dict(width=0.4, edgecolor='w'),
                pctdistance=0.75
            )
            plt.setp(autotexts_t, size=9, weight="bold", color="white")
            
            ax_t.legend(wedges_t, df_transp_nf['transportadora'], title="Transportadoras", loc="upper center", bbox_to_anchor=(0.5, -0.05), ncol=1, fontsize=8)
            ax_t.axis('equal')
            plt.tight_layout()
            st.pyplot(fig_t)

            st.dataframe(df_transp_nf.rename(columns={'transportadora': 'Transportadora', 'qtd_nfs': 'NFs Transportadas (Qtd)'}), use_container_width=True)

        st.markdown("---")

        c_tb1, c_tb2 = st.columns(2)

        # RESUMO TRANSPORTADORAS FULL
        with c_tb1:
            st.subheader("🚚 Resumo por Transportadora")
            df_transp_full = df_amostra.groupby('transportadora').agg({
                'valor_nf': 'sum',
                'valor_frete': 'sum',
                'nf': 'count'
            }).reset_index().rename(columns={'nf': 'qtd_nfs'}).sort_values(by='valor_nf', ascending=False)

            df_transp_view = df_transp_full.copy()
            df_transp_view['valor_nf'] = df_transp_view['valor_nf'].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
            df_transp_view['valor_frete'] = df_transp_view['valor_frete'].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))

            st.dataframe(
                df_transp_view.rename(columns={
                    'transportadora': 'Transportadora',
                    'valor_nf': 'Total Faturado (R$)',
                    'valor_frete': 'Total Frete (R$)',
                    'qtd_nfs': 'NFs'
                }),
                use_container_width=True
            )

        # RESUMO CLIENTES DECRESCENTE
        with c_tb2:
            st.subheader("🏢 Resumo por Cliente (Ordem Decrescente de Valor)")
            df_cli_full = df_amostra.groupby('cliente').agg({
                'valor_nf': 'sum',
                'valor_frete': 'sum',
                'nf': 'count'
            }).reset_index().rename(columns={'nf': 'qtd_nfs'}).sort_values(by='valor_nf', ascending=False)

            df_cli_view = df_cli_full.copy()
            df_cli_view['valor_nf'] = df_cli_view['valor_nf'].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
            df_cli_view['valor_frete'] = df_cli_view['valor_frete'].apply(lambda x: f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))

            st.dataframe(
                df_cli_view.rename(columns={
                    'cliente': 'Cliente / Destinatário',
                    'valor_nf': 'Total Faturado (R$)',
                    'valor_frete': 'Total Frete (R$)',
                    'qtd_nfs': 'NFs'
                }),
                use_container_width=True
            )
    else:
        st.warning("Nenhum registro encontrado para exibição da amostragem.")